"""
Doctor Directory API Routes for Arovida PET-CT Centre.
Handles contact management, bulk CSV/Excel imports, phone normalization (+91),
active/paused toggles, and sample CSV downloads.
"""
import io
import logging
from typing import List, Optional
import csv

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth import get_current_user
from app.database import DoctorContact, User, get_db
from app.schemas import (
    DoctorBulkUploadResult,
    DoctorCreate,
    DoctorResponse,
    DoctorUpdate,
    normalize_phone_number,
)

logger = logging.getLogger("whatsapp_ai.routes.doctors")
router = APIRouter(prefix="/api/doctors", tags=["Doctor Directory"])


@router.get("", response_model=List[DoctorResponse], summary="List doctor directory contacts")
async def list_doctors(
    q: Optional[str] = Query(None, description="Search query by name, phone, specialty, or hospital"),
    active_only: Optional[bool] = Query(None, description="Filter active doctors only"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    List all referring doctor contacts belonging to the authenticated centre.
    Supports real-time search across doctor name, hospital, and oncology specialty.
    """
    stmt = select(DoctorContact).where(DoctorContact.user_id == current_user.id)

    if active_only is True:
        stmt = stmt.where(DoctorContact.is_active == True)
    elif active_only is False:
        stmt = stmt.where(DoctorContact.is_active == False)

    if q and q.strip():
        term = f"%{q.strip()}%"
        stmt = stmt.where(
            or_(
                DoctorContact.doctor_name.ilike(term),
                DoctorContact.phone_number.ilike(term),
                DoctorContact.specialty.ilike(term),
                DoctorContact.hospital.ilike(term),
            )
        )

    stmt = stmt.order_by(DoctorContact.doctor_name.asc())
    result = await db.execute(stmt)
    return result.scalars().all()


@router.post("", response_model=DoctorResponse, status_code=status.HTTP_201_CREATED, summary="Add single doctor contact")
async def create_doctor(
    payload: DoctorCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Add a single referring doctor with automatic phone number sanitation (+91 default).
    """
    clean_phone = normalize_phone_number(payload.phone_number)

    # Check for existing doctor with same phone for this user
    existing_stmt = select(DoctorContact).where(
        DoctorContact.user_id == current_user.id,
        DoctorContact.phone_number == clean_phone,
    )
    existing_res = await db.execute(existing_stmt)
    if existing_res.scalar_one_or_none():
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Doctor with phone {clean_phone} already exists in your directory.",
        )

    doc = DoctorContact(
        user_id=current_user.id,
        doctor_name=payload.doctor_name.strip(),
        phone_number=clean_phone,
        specialty=payload.specialty.strip() if payload.specialty else None,
        hospital=payload.hospital.strip() if payload.hospital else None,
        is_active=payload.is_active,
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)
    return doc


@router.post("/bulk-upload", response_model=DoctorBulkUploadResult, summary="Bulk import doctor contacts from CSV or Excel")
async def bulk_upload_doctors(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Bulk import 500+ doctor contacts via CSV or Excel (.xlsx).
    Expected columns: Doctor Name, Phone, Specialty, Hospital.
    Automatically prepends +91 to 10-digit phone numbers and skips duplicates.
    """
    filename = file.filename.lower() if file.filename else ""
    content = await file.read()

    rows = []

    if filename.endswith(".csv"):
        try:
            decoded = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(decoded))
            rows = list(reader)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse CSV file: {str(e)}")

    elif filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            headers = [str(cell.value or "").strip() for cell in ws[1]]
            for r in ws.iter_rows(min_row=2, values_only=True):
                if any(r):
                    rows.append(dict(zip(headers, r)))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(e)}")
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload a .csv or .xlsx file.")

    if not rows:
        raise HTTPException(status_code=400, detail="Uploaded file contains no rows.")

    inserted = 0
    skipped = 0
    errors = []

    # Pre-fetch existing doctor phone numbers for this centre to avoid duplicate clashes
    existing_phones_res = await db.execute(
        select(DoctorContact.phone_number).where(DoctorContact.user_id == current_user.id)
    )
    existing_phones = set(existing_phones_res.scalars().all())

    def get_field(row_dict, candidates):
        for c in candidates:
            for k in row_dict.keys():
                if k and k.strip().lower() == c.lower():
                    val = row_dict[k]
                    return str(val).strip() if val is not None else ""
        return ""

    for idx, r in enumerate(rows, start=2):
        doc_name = get_field(r, ["Doctor Name", "Doctor", "Name", "Dr Name", "Full Name"])
        phone = get_field(r, ["Phone", "Phone Number", "Mobile", "Contact", "WhatsApp"])
        specialty = get_field(r, ["Specialty", "Specialization", "Department", "Field"])
        hospital = get_field(r, ["Hospital", "Affiliation", "Clinic", "Institution"])

        if not doc_name or not phone:
            errors.append(f"Row {idx}: Missing required Doctor Name or Phone Number.")
            continue

        try:
            clean_phone = normalize_phone_number(phone)
        except Exception as exc:
            errors.append(f"Row {idx} ({doc_name}): {exc}")
            continue

        if clean_phone in existing_phones:
            skipped += 1
            continue

        doc = DoctorContact(
            user_id=current_user.id,
            doctor_name=doc_name,
            phone_number=clean_phone,
            specialty=specialty if specialty else None,
            hospital=hospital if hospital else None,
            is_active=True,
        )
        db.add(doc)
        existing_phones.add(clean_phone)
        inserted += 1

    await db.commit()
    logger.info(f"Bulk upload complete for user [{current_user.id}]: {inserted} inserted, {skipped} skipped.")

    return DoctorBulkUploadResult(
        total_processed=len(rows),
        inserted=inserted,
        skipped_duplicates=skipped,
        errors=errors[:10],
    )


@router.get("/sample-csv", summary="Download sample Doctor Directory CSV template")
async def download_sample_csv():
    """
    Return a pre-formatted sample CSV with columns: Doctor Name, Phone, Specialty, Hospital.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Doctor Name", "Phone", "Specialty", "Hospital"])
    writer.writerow(["Dr. Rajesh Sharma", "9876543210", "Medical Oncology", "Tata Memorial Hospital"])
    writer.writerow(["Dr. Ananya Iyer", "9812345678", "Surgical Oncology", "Apollo Cancer Centre"])
    writer.writerow(["Dr. Vikram Malhotra", "9988776655", "Pulmonology & Chest Medicine", "Fortis Memorial Institute"])
    writer.writerow(["Dr. Sneha Patel", "9765432109", "Radiation Oncology", "Max Super Speciality Hospital"])
    writer.writerow(["Dr. Amit Verma", "9823456789", "Internal Medicine & Oncology", "City Care Hospital"])

    csv_data = output.getvalue()
    return Response(
        content=csv_data,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="Arovida_Doctor_Directory_Sample.csv"'},
    )


@router.patch("/{doctor_id}/toggle", response_model=DoctorResponse, summary="Toggle doctor active status")
async def toggle_doctor_status(
    doctor_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Toggle doctor status between Active (included in broadcasts) and Paused (excluded).
    """
    stmt = select(DoctorContact).where(
        DoctorContact.id == doctor_id,
        DoctorContact.user_id == current_user.id,
    )
    res = await db.execute(stmt)
    doc = res.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Doctor contact not found.")

    doc.is_active = not doc.is_active
    await db.commit()
    await db.refresh(doc)
    return doc


@router.delete("/{doctor_id}", status_code=status.HTTP_204_NO_CONTENT, summary="Delete doctor contact")
async def delete_doctor(
    doctor_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Delete a doctor contact from the directory.
    """
    stmt = select(DoctorContact).where(
        DoctorContact.id == doctor_id,
        DoctorContact.user_id == current_user.id,
    )
    res = await db.execute(stmt)
    doc = res.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Doctor contact not found.")

    await db.delete(doc)
    await db.commit()
    return None


@router.post("/seed-demo", summary="Seed 10 realistic referring doctors for testing")
async def seed_demo_doctors(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Instantly seed 10 realistic oncology & physician contacts for immediate demonstration.
    """
    demo_list = [
        ("Dr. Rajesh Sharma", "+919876543210", "Medical Oncology", "Tata Memorial Hospital"),
        ("Dr. Ananya Iyer", "+919812345678", "Surgical Oncology", "Apollo Cancer Centre"),
        ("Dr. Vikram Malhotra", "+919988776655", "Pulmonology & Chest Medicine", "Fortis Memorial Institute"),
        ("Dr. Sneha Patel", "+919765432109", "Radiation Oncology", "Max Super Speciality Hospital"),
        ("Dr. Amit Verma", "+919823456789", "Internal Medicine & Oncology", "City Care Hospital"),
        ("Dr. Meera Nambiar", "+919834567890", "Hemato-Oncology", "AIIMS Oncology Unit"),
        ("Dr. Siddharth Rao", "+919845678901", "Gastrointestinal Oncology", "Manipal Hospital"),
        ("Dr. Pooja Deshmukh", "+919856789012", "Urologic Oncology", "Kokilaben Dhirubhai Ambani Hospital"),
        ("Dr. Rohan Kulkarni", "+919867890123", "Neuro-Oncology", "Hinduja Hospital"),
        ("Dr. Kavita Singhania", "+919878901234", "Head & Neck Surgical Oncology", "Sir H. N. Reliance Hospital"),
    ]

    inserted = 0
    for name, phone, specialty, hospital in demo_list:
        existing = await db.execute(
            select(DoctorContact).where(
                DoctorContact.user_id == current_user.id,
                DoctorContact.phone_number == phone,
            )
        )
        if not existing.scalar_one_or_none():
            doc = DoctorContact(
                user_id=current_user.id,
                doctor_name=name,
                phone_number=phone,
                specialty=specialty,
                hospital=hospital,
                is_active=True,
            )
            db.add(doc)
            inserted += 1

    await db.commit()
    return {"message": f"Successfully loaded {inserted} demo doctor contacts into directory."}
